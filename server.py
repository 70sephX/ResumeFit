"""ResumeFit: small full-stack capstone MVP, stdlib SQLite, FastAPI, OpenAI/Stripe REST.
Note: production deployments must configure a persistent DB volume; ephemeral disks lose histories.
"""
from __future__ import annotations

import hashlib
import hmac
import json
import os
import re
import secrets
import sqlite3
import time
from datetime import datetime, timezone
from io import BytesIO
import zipfile
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request as URLRequest, urlopen

from fastapi import FastAPI, Request, HTTPException, Response, UploadFile, File
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import agent as agent_mod
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor

try:  # python.org builds on macOS ship without root certificates; certifi fixes HTTPS to OpenAI/Stripe
    import certifi
    os.environ.setdefault('SSL_CERT_FILE', certifi.where())
except ImportError:
    pass
ROOT = Path(__file__).resolve().parent
# Minimal .env loader so the project can start without extra dependencies
ENV_FILE = ROOT / '.env'
if ENV_FILE.exists():
    for line in ENV_FILE.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if line and not line.startswith('#') and '=' in line:
            k, v = line.split('=', 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

DB_FILE = Path(os.getenv('DATABASE_PATH', str(ROOT / 'data' / 'resumefit.db')))
DB_FILE.parent.mkdir(parents=True, exist_ok=True)
MODEL = os.getenv('OPENAI_MODEL', 'gpt-4.1-mini')
COOKIE = 'resumefit_session'
SECURE_COOKIE = os.getenv('COOKIE_SECURE', 'false').lower() == 'true'
FREE_CREDITS = 3
PAYMENT_CREDITS = 5
PAYMENT_USD_CENTS = 499

app = FastAPI(title='ResumeFit', version='2.0.0')
app.mount('/static', StaticFiles(directory=str(ROOT / 'static')), name='static')


def db() -> sqlite3.Connection:
    c = sqlite3.connect(str(DB_FILE), timeout=15)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    return c


def init_db() -> None:
    with db() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY, email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL, salt TEXT NOT NULL,
            credits INTEGER NOT NULL DEFAULT 3, created_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS sessions (
            token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL,
            expires_at INTEGER NOT NULL, FOREIGN KEY(user_id) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS runs (
            id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL,
            resume TEXT NOT NULL, job_description TEXT NOT NULL,
            result_json TEXT NOT NULL, created_at TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS facts (
            id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL, text TEXT NOT NULL,
            created_at TEXT NOT NULL, FOREIGN KEY(user_id) REFERENCES users(id)
        );
        CREATE TABLE IF NOT EXISTS stripe_sessions (
            session_id TEXT PRIMARY KEY, user_id INTEGER NOT NULL,
            credits INTEGER NOT NULL, created_at TEXT NOT NULL
        );
        ''')


init_db()


@app.get('/')
def index():
    return FileResponse(ROOT / 'static' / 'index.html')


@app.get('/health')
def health():
    return {'ok': True}


@app.get('/api/status')
def status():
    return {
        'ai_configured': bool(os.getenv('OPENAI_API_KEY')),
        'stripe_configured': bool(os.getenv('STRIPE_SECRET_KEY')),
        'model': MODEL if os.getenv('OPENAI_API_KEY') else None,
        'free_credits': FREE_CREDITS,
        'credit_pack': {'credits': PAYMENT_CREDITS, 'price_usd': PAYMENT_USD_CENTS / 100},
    }


class Credentials(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=8, max_length=128)


class TailorRequest(BaseModel):
    resume: str = Field(min_length=25, max_length=20000)
    job_description: str = Field(min_length=25, max_length=20000)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def hash_password(password: str, salt: str) -> str:
    return hashlib.pbkdf2_hmac('sha256', password.encode(), bytes.fromhex(salt), 230000).hex()


def get_user(request: Request) -> dict | None:
    token = request.cookies.get(COOKIE, '')
    if not token:
        return None
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    with db() as c:
        row = c.execute('''SELECT u.id, u.email, u.credits FROM sessions s
              JOIN users u ON u.id=s.user_id
              WHERE s.token_hash=? AND s.expires_at>?''', (token_hash, int(time.time()))).fetchone()
    return dict(row) if row else None


def require_user(request: Request) -> dict:
    user = get_user(request)
    if user is None:
        raise HTTPException(401, 'Please sign in to use this feature.')
    return user


def login_response(user_id: int) -> JSONResponse:
    token = secrets.token_urlsafe(36)
    with db() as c:
        c.execute('INSERT INTO sessions(token_hash,user_id,expires_at) VALUES(?,?,?)',
                  (hashlib.sha256(token.encode()).hexdigest(), user_id, int(time.time()) + 30 * 86400))
        row = c.execute('SELECT id,email,credits FROM users WHERE id=?', (user_id,)).fetchone()
    resp = JSONResponse({'user': dict(row)})
    resp.set_cookie(COOKIE, token, max_age=30 * 86400, httponly=True,
                    samesite='lax', secure=SECURE_COOKIE, path='/')
    return resp


@app.post('/api/auth/register')
def register(body: Credentials, request: Request):
    rate_limit('auth:' + (request.client.host if request.client else 'anon'), 10)
    email = body.email.strip().lower()
    if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
        raise HTTPException(400, 'Please enter a valid email address.')
    salt = secrets.token_hex(16)
    try:
        with db() as c:
            cur = c.execute('INSERT INTO users(email,password_hash,salt,credits,created_at) VALUES(?,?,?,?,?)',
                            (email, hash_password(body.password, salt), salt, FREE_CREDITS, now_iso()))
            user_id = cur.lastrowid
    except sqlite3.IntegrityError:
        raise HTTPException(409, 'An account with this email already exists.')
    return login_response(user_id)


@app.post('/api/auth/login')
def login(body: Credentials, request: Request):
    rate_limit('auth:' + (request.client.host if request.client else 'anon'), 10)
    with db() as c:
        row = c.execute('SELECT * FROM users WHERE email=?', (body.email.strip().lower(),)).fetchone()
    if not row or not hmac.compare_digest(hash_password(body.password, row['salt']), row['password_hash']):
        raise HTTPException(401, 'Email or password is incorrect.')
    return login_response(row['id'])


@app.post('/api/auth/logout')
def logout(request: Request):
    token = request.cookies.get(COOKIE, '')
    if token:
        with db() as c:
            c.execute('DELETE FROM sessions WHERE token_hash=?', (hashlib.sha256(token.encode()).hexdigest(),))
    resp = JSONResponse({'ok': True})
    resp.delete_cookie(COOKIE, path='/')
    return resp


@app.get('/api/me')
def me(request: Request):
    return {'user': get_user(request)}


def clean_text(s: str) -> str:
    return re.sub(r'\s+', ' ', s or '').strip()


def get_bullets(resume: str) -> list[str]:
    output: list[str] = []
    for line in resume.splitlines():
        l = line.strip()
        match = re.match(r'^(?:[-•*]|\d+[.)])\s*(.{16,})$', l)
        if match:
            output.append(match.group(1).strip())
    if not output:
        output = [clean_text(l) for l in resume.splitlines() if 35 < len(l.strip()) < 250][:5]
    return output[:12]


STOPWORDS = set('and the with for of in to on a an or as at by is are be we our you your this that from have has will do it they their across using use experience ability knowledge strong relevant excellent who skills skill work working plus including good understanding familiar such'.split())


def requirements_heuristic(jd: str) -> list[dict]:
    lines = [clean_text(x).lstrip('-•* 0123456789.)') for x in jd.splitlines()]
    lines = [x for x in lines if 16 <= len(x) <= 180 and len(x.split()) >= 3]
    needles = ['python', 'sql', 'data analysis', 'machine learning', 'visualization', 'collaboration', 'communication', 'react', 'tableau', 'statistics', 'marketing', 'product', 'analytics', 'excel']
    items: list[str] = []
    for line in lines:
        if any(w in line.lower() for w in needles) or re.search(r'\b(required|develop|manage|analyze|build|design|collaborate|communicate|experience|proficien|responsib)\b', line, re.I):
            items.append(line.rstrip('.'))
    if len(items) < 3:
        items = lines
    return [{'text': x[:140], 'status': 'not_evidenced', 'evidence': ''} for x in items[:6]]


def safe_demo(resume: str, jd: str) -> dict:
    """A transparent rule-based preview, not an AI model."""
    lower_resume = resume.lower()
    requirements = requirements_heuristic(jd)
    for r in requirements:
        words = [w for w in re.findall(r'[a-z][a-z0-9+#.]{2,}', r['text'].lower()) if w not in STOPWORDS]
        found = [word for word in words if re.search(r'\b' + re.escape(word) + r'\b', lower_resume)]
        r['status'] = 'matched' if len(found) >= 2 else 'partial' if found else 'not_evidenced'
        # Show a verbatim supporting resume sentence, not just keyword fragments.
        support = ''
        for bullet in get_bullets(resume):
            bullet_hits = sum(1 for word in words if re.search(r'\b' + re.escape(word) + r'\b', bullet, re.I))
            if bullet_hits >= (2 if r['status'] == 'matched' else 1):
                support = bullet
                break
        if not support and found:
            support = next((line.strip() for line in resume.splitlines() if any(re.search(r'\b'+re.escape(w)+r'\b',line,re.I) for w in found)), '')
        r['evidence'] = support[:250] if found else ''
    raw_bullets = get_bullets(resume)
    suggestions = []
    rewrites = {
        'worked on': 'Contributed to', 'helped with': 'Supported',
        'was responsible for': 'Managed', 'created': 'Developed',
        'made': 'Created', 'did': 'Completed', 'used': 'Applied',
        'built': 'Developed', 'assisted with': 'Supported',
    }
    for original in raw_bullets[:5]:
        after = original
        for prefix, replacement in rewrites.items():
            if original.lower().startswith(prefix + ' '):
                after = replacement + original[len(prefix):]
                break
        if after == original:
            after = original[0].upper() + original[1:].rstrip('.') + '.'
        elif after[-1] != '.':
            after += '.'
        suggestions.append({
            'before': original, 'after': after,
            'why': 'Uses clearer action-oriented phrasing while retaining the original claim. Review before using.',
            'evidence': original,
        })
    if not suggestions:
        suggestions = []
    missing = [r['text'] for r in requirements if r['status'] == 'not_evidenced'][:5]
    return {
        'requirements': requirements, 'suggestions': suggestions,
        'missing': missing,
        'summary': 'This is a rule-based demo preview. Add OPENAI_API_KEY to enable context-aware AI rewriting and analysis. No claims should be added without resume evidence.',
        'mode': 'demo',
    }


SYSTEM_PROMPT = '''You are ResumeFit, a careful resume editor. Tailor the user's resume to the target job WITHOUT INVENTING ANY EXPERIENCES. Return ONLY a JSON object matching the keys below.
Never fabricate or infer a skill, tool, team, job title, responsibility, accomplishment, employer, outcome or metric that the resume does not explicitly support. A job description is NOT evidence that the person has a skill. Reorder and rephrase existing facts only. Do not make up percentages. Do not promise a score or interview outcome.
Find 3-6 short JD requirements. Each requirement is {"text": str, "status": "matched"|"partial"|"not_evidenced", "evidence": str}. Evidence should quote the actual resume text if matched. If not evidenced, set evidence to "". Use "partial" where indirect evidence may exist; do not overclaim.
Create 3-5 rewritten bullet suggestions. Each suggestion is {"before": str, "after": str, "why": str, "evidence": str}. The "before" MUST be an EXACT contiguous substring from the resume (prefer a full bullet without bullet marker). The "evidence" MUST be an exact contiguous substring from the resume, supporting all factual claims. "after" must contain ONLY facts explicit in before/evidence, while tailoring focus to JD. Short impactful bullets. If resume has fewer than 3 suitable statements, produce fewer rather than inventing.
"missing": array of 0-5 concise JD requirements not clearly evidenced in resume (no assumption of lack of ability). "summary": 1-2 short sentences of actionable guidance, based strictly on supplied content.
Output the JSON object only: {"requirements": [...], "suggestions": [...], "missing": [...], "summary": "..."}.'''


def call_openai(resume: str, jd: str) -> dict:
    key = os.getenv('OPENAI_API_KEY', '')
    body = {
        'model': MODEL,
        'temperature': 0.2,
        'response_format': {'type': 'json_object'},
        'messages': [
            {'role': 'system', 'content': SYSTEM_PROMPT},
            {'role': 'user', 'content': f'<resume>\n{resume}\n</resume>\n\n<job_description>\n{jd}\n</job_description>'},
        ],
    }
    req = URLRequest('https://api.openai.com/v1/chat/completions',
                  data=json.dumps(body).encode(),
                  headers={'Authorization': f'Bearer {key}', 'Content-Type': 'application/json'}, method='POST')
    try:
        with urlopen(req, timeout=55) as r:
            data = json.loads(r.read().decode())
        result = json.loads(data['choices'][0]['message']['content'])
    except HTTPError as e:
        raise HTTPException(502, f'AI provider error ({e.code}). Check your API key and model access.')
    except TimeoutError:
        raise HTTPException(504, 'AI provider timed out. Please try again.')
    except URLError as e:
        raise HTTPException(502, f'Could not reach the AI provider: {e.reason}')
    except (ValueError, KeyError, IndexError, TypeError):
        raise HTTPException(502, 'AI did not return a valid structured response.')
    return validate_ai(result, resume)


def validate_ai(result: dict, resume: str) -> dict:
    if not isinstance(result, dict):
        raise HTTPException(502, 'Invalid AI result shape.')
    norm_resume = clean_text(resume).casefold()
    reqs = []
    for item in result.get('requirements', [])[:8]:
        if not isinstance(item, dict):
            continue
        label = clean_text(str(item.get('text', '')))[:170]
        evidence = clean_text(str(item.get('evidence', '')))[:300]
        if evidence and evidence.casefold() not in norm_resume:
            evidence = ''
        if not label:
            continue
        status = item.get('status', 'not_evidenced')
        if status not in ('matched', 'partial', 'not_evidenced'):
            status = 'not_evidenced'
        if not evidence and status == 'matched':
            status = 'not_evidenced'
        reqs.append({'text': label, 'status': status, 'evidence': evidence})
    suggestions = []
    for item in result.get('suggestions', [])[:5]:
        if not isinstance(item, dict):
            continue
        before = clean_text(str(item.get('before', '')))[:650]
        after = clean_text(str(item.get('after', '')))[:650]
        evidence = clean_text(str(item.get('evidence', '')))[:650]
        if not before or not after or before.casefold() not in norm_resume or evidence.casefold() not in norm_resume or len(evidence) < 12:
            continue
        # Reject novel numeric claims. Human review still required for unsupported non-numeric claims.
        nums = re.findall(r'(?<!\w)\d+(?:[.,]\d+)?%?', after)
        resume_nums = re.findall(r'(?<!\w)\d+(?:[.,]\d+)?%?', resume)
        if any(n not in resume_nums for n in nums):
            continue
        # Lightweight additional safeguard: reject introducing named tools absent
        # from the source. It supplements, but cannot replace, human fact checking.
        named_tools = ['Python', 'SQL', 'Tableau', 'Excel', 'Power BI', 'AWS', 'GCP',
                       'Azure', 'React', 'Node.js', 'MongoDB', 'Docker', 'Kubernetes',
                       'PyTorch', 'TensorFlow', 'Spark', 'BigQuery', 'Databricks',
                       'Snowflake', 'SAS', 'MATLAB', 'JavaScript', 'GitHub']
        def mentions_tool(text: str, tool: str) -> bool:
            return bool(re.search(r'(?<!\w)' + re.escape(tool) + r'(?!\w)', text, re.I))
        if any(mentions_tool(after, tool) and not mentions_tool(resume, tool)
               for tool in named_tools):
            continue
        suggestions.append({'before': before, 'after': after, 'why': clean_text(str(item.get('why', '')))[:400], 'evidence': evidence})
    if not reqs:
        reqs = requirements_heuristic('')
    missing = [clean_text(str(x))[:180] for x in result.get('missing', [])[:5] if isinstance(x, str)]
    return {'requirements': reqs, 'suggestions': suggestions, 'missing': missing,
            'summary': clean_text(str(result.get('summary', '')))[:1000], 'mode': 'ai'}


def run_ai_tailor(resume: str, jd: str, bank: list[str]) -> dict:
    """Run the agent; translate agent errors into HTTP errors (fail loud, never a fake answer)."""
    usage: dict = {}
    try:
        result = agent_mod.run_agent(resume, jd, agent_mod.openai_llm(usage), bank)
    except agent_mod.AgentError as e:
        raise HTTPException(e.status, e.message)
    result['agent']['tokens'] = usage
    return result


RATE: dict[str, list[float]] = {}


def rate_limit(key: str, limit: int, window: int = 60) -> None:
    """Tiny in-memory sliding window (per process). Enough for a demo; use Redis when scaled out."""
    now = time.time()
    hits = [t for t in RATE.get(key, []) if now - t < window]
    if len(hits) >= limit:
        raise HTTPException(429, 'Too many requests. Please wait a minute and try again.')
    RATE[key] = hits + [now]


@app.post('/api/tailor')
def tailor(body: TailorRequest, request: Request):
    resume = body.resume.strip()
    jd = body.job_description.strip()
    user = get_user(request)
    rate_limit('tailor:' + (str(user['id']) if user else (request.client.host if request.client else 'anon')), 10)
    ai_on = bool(os.getenv('OPENAI_API_KEY'))
    if ai_on and user is None:
        raise HTTPException(401, 'Create a free account to use AI tailoring.')
    if ai_on and user['credits'] < 1:
        raise HTTPException(402, 'You have no remaining credits. Purchase more to continue.')
    if ai_on:
        with db() as c:
            bank = [r['text'] for r in c.execute('SELECT text FROM facts WHERE user_id=? ORDER BY id DESC LIMIT 60', (user['id'],))]
        result = run_ai_tailor(resume, jd, bank)
    else:
        result = safe_demo(resume, jd)
    # Charge only for a grounded result: nothing is deducted on errors (raised above) or when
    # the verifier left nothing usable (no verified rewrite and no requirement with evidence).
    grounded = bool(result['suggestions']) or any(r['status'] != 'not_evidenced' for r in result['requirements'])
    charge = ai_on and grounded
    result['charged'] = charge
    if user:
        with db() as c:
            c.execute('BEGIN IMMEDIATE')
            if charge:
                cur = c.execute('UPDATE users SET credits=credits-1 WHERE id=? AND credits>0', (user['id'],))
                if cur.rowcount != 1:
                    raise HTTPException(402, 'You have no remaining credits.')
            row = c.execute('INSERT INTO runs(user_id,resume,job_description,result_json,created_at) VALUES(?,?,?,?,?)',
                            (user['id'], resume, jd, json.dumps(result), now_iso()))
            result['run_id'] = row.lastrowid
            result['credits_remaining'] = c.execute('SELECT credits FROM users WHERE id=?', (user['id'],)).fetchone()[0]
    return result


class FactIn(BaseModel):
    text: str = Field(min_length=14, max_length=400)


@app.get('/api/facts')
def list_facts(request: Request):
    user = require_user(request)
    with db() as c:
        rows = c.execute('SELECT id,text,created_at FROM facts WHERE user_id=? ORDER BY id DESC', (user['id'],)).fetchall()
    return {'items': [dict(r) for r in rows]}


@app.post('/api/facts')
def add_fact(body: FactIn, request: Request):
    user = require_user(request)
    text = clean_text(body.text)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        if c.execute('SELECT COUNT(*) FROM facts WHERE user_id=?', (user['id'],)).fetchone()[0] >= 60:
            raise HTTPException(409, 'Fact bank is full (60 entries). Delete some before adding more.')
        if c.execute('SELECT 1 FROM facts WHERE user_id=? AND lower(text)=lower(?)', (user['id'], text)).fetchone():
            raise HTTPException(409, 'That fact is already saved.')
        row = c.execute('INSERT INTO facts(user_id,text,created_at) VALUES(?,?,?)', (user['id'], text, now_iso()))
    return {'id': row.lastrowid, 'text': text}


@app.delete('/api/facts/{fact_id}')
def delete_fact(fact_id: int, request: Request):
    user = require_user(request)
    with db() as c:
        cur = c.execute('DELETE FROM facts WHERE id=? AND user_id=?', (fact_id, user['id']))
    if cur.rowcount != 1:
        raise HTTPException(404, 'Fact not found.')
    return {'deleted': True}


@app.get('/api/history')
def history(request: Request):
    user = require_user(request)
    with db() as c:
        items = c.execute('''SELECT id,job_description,created_at, result_json FROM runs
              WHERE user_id=? ORDER BY id DESC LIMIT 50''', (user['id'],)).fetchall()
    output = []
    for item in items:
        title = next((clean_text(x)[:65] for x in item['job_description'].splitlines() if x.strip()), 'Job description')
        obj = json.loads(item['result_json'])
        output.append({'id': item['id'], 'title': title, 'created_at': item['created_at'], 'mode': obj.get('mode', 'demo')})
    return {'items': output}


@app.get('/api/history/{run_id}')
def history_detail(run_id: int, request: Request):
    user = require_user(request)
    with db() as c:
        row = c.execute('SELECT * FROM runs WHERE id=? AND user_id=?', (run_id, user['id'])).fetchone()
    if row is None:
        raise HTTPException(404, 'History entry not found.')
    return {'id': row['id'], 'resume': row['resume'], 'job_description': row['job_description'],
            'result': json.loads(row['result_json']), 'created_at': row['created_at']}


@app.post('/api/billing/checkout')
def checkout(request: Request):
    user = require_user(request)
    key = os.getenv('STRIPE_SECRET_KEY', '')
    if not key:
        raise HTTPException(503, 'Stripe test checkout is not configured yet.')
    origin = os.getenv('APP_BASE_URL', '').rstrip('/')
    if not origin:
        raise HTTPException(503, 'Set APP_BASE_URL before enabling Stripe.')
    params = {
        'mode': 'payment', 'client_reference_id': str(user['id']),
        'success_url': origin + '/?payment=success',
        'cancel_url': origin + '/?payment=cancel',
        'line_items[0][price_data][currency]': 'usd',
        'line_items[0][price_data][product_data][name]': 'ResumeFit – 5 AI credits',
        'line_items[0][price_data][unit_amount]': str(PAYMENT_USD_CENTS),
        'line_items[0][quantity]': '1',
        'metadata[user_id]': str(user['id']), 'metadata[credits]': str(PAYMENT_CREDITS),
    }
    req = URLRequest('https://api.stripe.com/v1/checkout/sessions',
                  data=urlencode(params).encode(), method='POST',
                  headers={'Authorization': 'Bearer ' + key, 'Content-Type': 'application/x-www-form-urlencoded'})
    try:
        with urlopen(req, timeout=20) as r:
            obj = json.loads(r.read())
        return {'url': obj['url']}
    except HTTPError as e:
        raise HTTPException(502, f'Stripe checkout error ({e.code}).')
    except (URLError, ValueError, KeyError):
        raise HTTPException(502, 'Could not create checkout session.')


def verify_stripe_signature(payload: bytes, header: str, secret: str) -> bool:
    vals = {}
    for part in header.split(','):
        if '=' in part:
            key, value = part.strip().split('=', 1)
            vals.setdefault(key, []).append(value)
    if 't' not in vals or 'v1' not in vals:
        return False
    try:
        ts = int(vals['t'][0])
    except ValueError:
        return False
    if abs(time.time() - ts) > 300:
        return False
    expected = hmac.new(secret.encode(), str(ts).encode() + b'.' + payload, hashlib.sha256).hexdigest()
    return any(hmac.compare_digest(expected, v) for v in vals['v1'])


@app.post('/api/billing/webhook')
async def stripe_webhook(request: Request):
    secret = os.getenv('STRIPE_WEBHOOK_SECRET', '')
    payload = await request.body()
    if not secret or not verify_stripe_signature(payload, request.headers.get('stripe-signature', ''), secret):
        raise HTTPException(400, 'Invalid Stripe signature.')
    event = json.loads(payload)
    if event.get('type') == 'checkout.session.completed':
        obj = event['data']['object']
        if obj.get('payment_status') == 'paid' and obj.get('mode') == 'payment':
            session_id = str(obj.get('id', ''))
            try:
                uid = int(obj.get('metadata', {}).get('user_id', ''))
            except ValueError:
                return {'received': True}
            # Grant only fixed official package and only once; verify metadata + paid amount.
            if (obj.get('currency') == 'usd' and obj.get('amount_total') == PAYMENT_USD_CENTS
                    and obj.get('metadata', {}).get('credits') == str(PAYMENT_CREDITS)):
                with db() as c:
                    c.execute('BEGIN IMMEDIATE')
                    row = c.execute('INSERT OR IGNORE INTO stripe_sessions(session_id,user_id,credits,created_at) VALUES(?,?,?,?)',
                                    (session_id, uid, PAYMENT_CREDITS, now_iso()))
                    if row.rowcount:
                        c.execute('UPDATE users SET credits=credits+? WHERE id=?', (PAYMENT_CREDITS, uid))
    return {'received': True}



# Uploads are parsed in memory; resume files are never written to disk.
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 20000


def extract_document_bytes(filename: str, raw: bytes) -> tuple[str, str]:
    """Extract plain text from a small PDF/DOCX/TXT safely and transparently."""
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.pdf', '.docx', '.txt'):
        raise HTTPException(415, 'Please upload a PDF, Word .docx, or .txt file.')
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise HTTPException(413, 'File must be smaller than 5 MB and not empty.')
    try:
        if suffix == '.pdf':
            if not raw.startswith(b'%PDF'):
                raise HTTPException(422, 'This file is not a valid PDF.')
            reader = PdfReader(BytesIO(raw), strict=False)
            if reader.is_encrypted:
                raise HTTPException(422, 'Password-protected PDFs are not supported. Please unlock your file first.')
            if len(reader.pages) > 20:
                raise HTTPException(413, 'Please upload a PDF containing no more than 20 pages.')
            text = '\n\n'.join((page.extract_text() or '') for page in reader.pages)
        elif suffix == '.docx':
            if not zipfile.is_zipfile(BytesIO(raw)):
                raise HTTPException(422, 'This file is not a valid .docx document.')
            with zipfile.ZipFile(BytesIO(raw)) as archive:
                if len(archive.namelist()) > 250 or sum(f.file_size for f in archive.infolist()) > 18 * 1024 * 1024:
                    raise HTTPException(413, 'The Word document is too complex or large to process.')
            doc = Document(BytesIO(raw))
            lines = [para.text for para in doc.paragraphs if para.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            text = '\n'.join(lines)
        else:
            text = raw.decode('utf-8-sig')
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, 'Could not read the document. Please try another export or paste the text.') from exc
    text = text.replace('\x00', '').strip()
    if len(text) < 25:
        raise HTTPException(422, 'No readable text was found. Scanned/image-only PDFs need OCR; please paste the resume text instead.')
    if len(text) > MAX_EXTRACTED_CHARS:
        raise HTTPException(413, 'The extracted text exceeds 20,000 characters. Shorten the document and retry.')
    return text, suffix.lstrip('.')


@app.post('/api/upload')
async def upload_document(file: UploadFile = File(...)):
    try:
        data = await file.read(MAX_FILE_BYTES + 1)
        text, file_type = extract_document_bytes(file.filename or '', data)
        return {'text': text, 'filename': Path(file.filename or '').name[:120],
                'file_type': file_type, 'characters': len(text)}
    finally:
        await file.close()


class ExportRequest(BaseModel):
    resume: str = Field(min_length=25, max_length=20000)
    title: str = Field(default='Tailored Resume', max_length=100)


@app.post('/api/export/docx')
def export_resume_docx(body: ExportRequest):
    """Export reviewed, edited plain text as an ATS-friendly single-column Word document."""
    doc = Document()
    sect = doc.sections[0]
    sect.top_margin = sect.bottom_margin = Inches(.73)
    sect.left_margin = sect.right_margin = Inches(.83)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(28, 42, 67)
    normal.paragraph_format.space_after = Pt(4)
    for idx, line in enumerate(body.resume.splitlines()):
        content = line.strip()
        if not content:
            continue
        paragraph = doc.add_paragraph()
        if idx == 0:
            run = paragraph.add_run(content)
            run.bold = True
            run.font.size = Pt(17)
            run.font.color.rgb = RGBColor(19, 49, 94)
        elif content.isupper() and len(content) < 65 and len(content) >= 3:
            run = paragraph.add_run(content)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(26, 79, 168)
            paragraph.paragraph_format.space_before = Pt(10)
        else:
            if content.startswith(('• ', '- ', '* ')):
                content = '• ' + content[2:].strip()
            paragraph.add_run(content)
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            headers={'Content-Disposition': 'attachment; filename="ResumeFit_Tailored_Resume.docx"', 'Cache-Control':'no-store'})


@app.get('/api/showcase')
def showcase():
    """Reproducible, no-key, honest showcase: never consumes real-AI credits."""
    example_resume = """JANE DOE
Data Analytics Student
EDUCATION
B.S. in Statistics and Data Science
EXPERIENCE
• Analyzed retail sales data using Python and SQL to identify purchasing patterns.
• Created Tableau dashboards to communicate weekly business metrics to a team of analysts.
• Built machine learning classification models for a university project.
• Collaborated with three classmates to present findings to a course instructor.
• Used Excel to clean and summarize marketing campaign data."""
    example_jd = """Junior Data Analyst — Growth Team
Requirements:
• Experience with Python and SQL for data analysis.
• Build clear dashboards and reports using Tableau or similar tools.
• Familiarity with machine learning fundamentals.
• Communicate insights to non-technical audiences.
• Experience with cloud platforms such as AWS or GCP.
• Collaborate across teams to support business decisions."""
    return {'resume': example_resume, 'job_description': example_jd,
            'result': safe_demo(example_resume, example_jd)}


@app.delete('/api/history/{run_id}')
def delete_history(run_id: int, request: Request):
    user = require_user(request)
    with db() as c:
        cursor = c.execute('DELETE FROM runs WHERE id=? AND user_id=?', (run_id, user['id']))
    if cursor.rowcount == 0:
        raise HTTPException(404, 'History entry not found.')
    return {'ok': True}


@app.delete('/api/account')
def delete_account(request: Request):
    user = require_user(request)
    with db() as c:
        c.execute('BEGIN IMMEDIATE')
        c.execute('DELETE FROM runs WHERE user_id=?', (user['id'],))
        c.execute('DELETE FROM facts WHERE user_id=?', (user['id'],))
        c.execute('DELETE FROM sessions WHERE user_id=?', (user['id'],))
        c.execute('DELETE FROM stripe_sessions WHERE user_id=?', (user['id'],))
        c.execute('DELETE FROM users WHERE id=?', (user['id'],))
    response = JSONResponse({'ok': True})
    response.delete_cookie(COOKIE, path='/')
    return response
