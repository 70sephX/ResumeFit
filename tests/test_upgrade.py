"""API + upload/export integration tests: no secrets or external calls needed."""
import io
import os
import unittest
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfWriter
from server import app, extract_document_bytes

class UpgradeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ.pop('OPENAI_API_KEY', None)  # these tests cover demo mode, whatever .env contains
        cls.client=TestClient(app)

    def test_showcase(self):
        r=self.client.get('/api/showcase')
        self.assertEqual(r.status_code,200)
        data=r.json()
        self.assertEqual(data['result']['mode'],'demo')
        self.assertTrue(len(data['result']['suggestions'])>=3)
        for row in data['result']['suggestions']:
            self.assertIn(row['evidence'],data['resume'])

    def test_docx_upload(self):
        doc=Document()
        doc.add_paragraph('JANE DOE')
        doc.add_paragraph('• Used Python and SQL for reporting.')
        file=io.BytesIO();doc.save(file)
        r=self.client.post('/api/upload',files={'file':('resume.docx',file.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')})
        self.assertEqual(r.status_code,200, r.text)
        self.assertIn('Python and SQL',r.json()['text'])

    def test_pdf_extraction(self):
        # Blank PDF has no selectable text; useful explanatory error, never fake OCR.
        writer=PdfWriter();writer.add_blank_page(width=612,height=792)
        buf=io.BytesIO();writer.write(buf)
        r=self.client.post('/api/upload',files={'file':('resume.pdf',buf.getvalue(),'application/pdf')})
        self.assertEqual(r.status_code,422)
        self.assertIn('No readable text',r.json()['detail'])

    def test_bad_files(self):
        r=self.client.post('/api/upload',files={'file':('resume.exe',b'not-a-document','application/octet-stream')})
        self.assertEqual(r.status_code,415)
        r=self.client.post('/api/upload',files={'file':('bad.pdf',b'not-a-pdf','application/pdf')})
        self.assertEqual(r.status_code,422)
        r=self.client.post('/api/upload',files={'file':('huge.txt',b'x'*(5*1024*1024+1),'text/plain')})
        self.assertEqual(r.status_code,413)

    def test_word_export_and_roundtrip(self):
        resume='JANE DOE\nEXPERIENCE\n• Used Python and SQL to analyze customer purchases and summarize monthly trends.'
        r=self.client.post('/api/export/docx',json={'resume':resume})
        self.assertEqual(r.status_code,200,r.text)
        self.assertIn('officedocument',r.headers['content-type'])
        text,kind=extract_document_bytes('export.docx',r.content)
        self.assertIn('customer purchases',text)
        self.assertEqual(kind,'docx')

    def test_demo_tailor_custom_input(self):
        r=self.client.post('/api/tailor',json={
            'resume':'EXPERIENCE\n• Used Python and SQL to analyze retail data for customer retention metrics.\n• Designed Tableau dashboard for metrics.',
            'job_description':'Data Analyst\n• Analyze data with Python and SQL.\n• Communicate with stakeholders using dashboards.\n• Build solutions using AWS cloud.'})
        self.assertEqual(r.status_code,200, r.text)
        self.assertEqual(r.json()['mode'],'demo')
        self.assertTrue(r.json()['suggestions'])

if __name__=='__main__':unittest.main()
